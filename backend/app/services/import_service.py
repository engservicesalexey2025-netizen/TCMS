"""Generic Excel import pipeline (Тома 1 §12 / Тома 1А §15 stages):

Выбор файла → Определение шаблона → Нормализация → Проверка → Сверка →
Показ различий → Подтверждение → Запись в PostgreSQL.

Implemented once against an ImportProfile (see import_profiles.py) — adding
import for a new entity does not touch this module.
"""
import uuid
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

import openpyxl
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models.enums import AuditAction, ImportRowAppliedAction as AppliedAction
from app.models.enums import ImportRowClassification as Classification
from app.models.enums import ImportRowDecision as Decision
from app.models.enums import ImportSessionStatus as Status
from app.models.import_session import ImportRow, ImportSession
from app.services.audit import log_action
from app.services.import_profiles import ImportProfile, get_profile

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 20_000

CONFIRMABLE_CLASSIFICATIONS = {Classification.NEW, Classification.IDENTICAL, Classification.CONFLICT}


class ImportRejected(Exception):
    """Raised when the whole file must be rejected before any row is staged."""


def _json_safe(value):
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, uuid.UUID):
        return str(value)
    if isinstance(value, Decimal):
        return float(value)
    if hasattr(value, "value") and hasattr(value, "name") and not isinstance(value, (int, str)):
        # enum.Enum member
        return value.value
    return value


def _json_safe_dict(d: dict | None) -> dict | None:
    if d is None:
        return None
    return {k: _json_safe(v) for k, v in d.items()}


def _rehydrate_raw(profile: ImportProfile, raw: dict) -> dict:
    """RawData уже прошла через _json_safe_dict при сохранении (дата стала
    ISO-строкой), поэтому перед повторной нормализацией нередактированных
    полей их нужно вернуть к date/datetime — иначе _normalize_row попытается
    распарсить ISO-строку форматом "%d.%m.%Y" из исходного файла и упадёт."""
    result = dict(raw)
    for col in profile.columns:
        value = result.get(col.field)
        if col.kind == "date" and isinstance(value, str):
            try:
                result[col.field] = date.fromisoformat(value.split("T")[0])
            except ValueError:
                pass
    return result


def _match_and_classify(
    db: Session, profile: ImportProfile, normalized: dict
) -> tuple[Classification, uuid.UUID | None, int | None, dict | None]:
    """Сверка нормализованной строки с текущими активными записями сущности —
    общая логика для первичного парсинга файла и повторной проверки строки
    после ручного исправления в разделе «Ошибки»."""
    key = tuple(normalized.get(f) for f in profile.match_fields)
    filters = [getattr(profile.model, f) == v for f, v in zip(profile.match_fields, key)]
    matches = db.scalars(
        select(profile.model).where(profile.model.IsActive.is_(True), *filters)
    ).all()
    if len(matches) == 0:
        return Classification.NEW, None, None, None
    if len(matches) == 1:
        obj = matches[0]
        diff = profile.repo.diff_fields(obj, normalized)
        if diff:
            diff_fields = {
                f: {"old": _json_safe(old), "new": _json_safe(new)}
                for f, (old, new) in diff.items()
            }
            return Classification.CONFLICT, obj.Id, obj.Version, diff_fields
        return Classification.IDENTICAL, obj.Id, obj.Version, None
    return Classification.AMBIGUOUS, None, None, None


def _duplicate_within_session(
    db: Session, row: ImportRow, profile: ImportProfile, normalized: dict
) -> bool:
    key = tuple(_json_safe(normalized.get(f)) for f in profile.match_fields)
    other_rows = db.scalars(
        select(ImportRow).where(
            ImportRow.ImportSessionId == row.ImportSessionId,
            ImportRow.Id != row.Id,
            ImportRow.Classification != Classification.ERROR,
        )
    ).all()
    for other in other_rows:
        if other.NormalizedData is None:
            continue
        other_key = tuple(other.NormalizedData.get(f) for f in profile.match_fields)
        if other_key == key:
            return True
    return False


def _normalize_header(value) -> str | None:
    if value is None:
        return None
    return str(value).strip().lower()


def _detect_columns(header_row, profile: ImportProfile) -> dict[str, int | None]:
    normalized_headers = [_normalize_header(h) for h in header_row]
    mapping: dict[str, int | None] = {}
    for col in profile.columns:
        aliases = {a.strip().lower() for a in col.header_aliases}
        idx = next((i for i, h in enumerate(normalized_headers) if h in aliases), None)
        if idx is None and col.required:
            raise ImportRejected(
                f"Обязательная колонка не найдена в файле: {col.header_aliases[0]}"
            )
        mapping[col.field] = idx
    return mapping


def _normalize_row(db: Session, profile: ImportProfile, raw: dict) -> tuple[dict | None, str | None]:
    normalized = {}
    for col in profile.columns:
        value = raw.get(col.field)
        if isinstance(value, str):
            value = value.strip() or None
        if value is None:
            if col.required:
                return None, f"Обязательное поле не заполнено: {col.field}"
            normalized[col.field] = None
            continue
        try:
            if col.kind == "text":
                normalized[col.field] = str(value)
            elif col.kind == "date":
                if isinstance(value, datetime):
                    normalized[col.field] = value.date()
                elif isinstance(value, date):
                    normalized[col.field] = value
                else:
                    normalized[col.field] = datetime.strptime(str(value).strip(), "%d.%m.%Y").date()
            elif col.kind == "enum":
                text = str(value).strip()
                mapped = col.enum_map.get(text, text) if col.enum_map else text
                if mapped not in col.enum_cls.__members__:
                    return None, f"Недопустимое значение поля {col.field}: '{value}'"
                normalized[col.field] = col.enum_cls[mapped]
            elif col.kind == "fk_lookup":
                text = str(value).strip()
                match = db.scalar(
                    select(col.fk_model).where(
                        getattr(col.fk_model, col.fk_lookup_field) == text,
                        col.fk_model.IsActive.is_(True),
                    )
                )
                if match is None:
                    return None, f"Справочник не найден ({col.field}): '{text}'"
                normalized[col.field] = match.Id
        except (ValueError, TypeError) as exc:
            return None, f"Ошибка в поле {col.field}: {exc}"
    return normalized, None


def parse_and_stage(
    db: Session,
    profile: ImportProfile,
    file_bytes: bytes,
    filename: str,
    user_id: uuid.UUID | None,
) -> ImportSession:
    """Stages 1-6: parse, detect template, normalize, validate, reconcile."""
    if not filename.lower().endswith(".xlsx"):
        session = ImportSession(
            EntityType=profile.entity_type, FileName=filename, Status=Status.FAILED,
            ErrorMessage="Поддерживаются только файлы .xlsx", CreatedBy=user_id,
        )
        db.add(session)
        db.commit()
        db.refresh(session)
        return session

    if len(file_bytes) > MAX_FILE_BYTES:
        session = ImportSession(
            EntityType=profile.entity_type, FileName=filename, Status=Status.FAILED,
            ErrorMessage=f"Файл превышает лимит {MAX_FILE_BYTES // (1024 * 1024)} МБ",
            CreatedBy=user_id,
        )
        db.add(session)
        db.commit()
        db.refresh(session)
        return session

    try:
        workbook = openpyxl.load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
        worksheet = workbook.active
        rows_iter = worksheet.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        if header_row is None:
            raise ImportRejected("Файл пуст")
        column_map = _detect_columns(header_row, profile)
    except ImportRejected as exc:
        session = ImportSession(
            EntityType=profile.entity_type, FileName=filename, Status=Status.FAILED,
            ErrorMessage=str(exc), CreatedBy=user_id,
        )
        db.add(session)
        db.commit()
        db.refresh(session)
        return session
    except Exception:
        session = ImportSession(
            EntityType=profile.entity_type, FileName=filename, Status=Status.FAILED,
            ErrorMessage="Не удалось прочитать файл — убедитесь, что это корректный .xlsx",
            CreatedBy=user_id,
        )
        db.add(session)
        db.commit()
        db.refresh(session)
        return session

    session = ImportSession(
        EntityType=profile.entity_type,
        FileName=filename,
        Status=Status.UPLOADED,
        ColumnMapping=dict(column_map),
        CreatedBy=user_id,
    )
    db.add(session)
    db.flush()

    raw_rows: list[tuple[int, dict]] = []
    for row_number, row in enumerate(rows_iter, start=2):
        if row is None or all(v is None for v in row):
            continue
        if len(raw_rows) >= MAX_ROWS:
            session.Status = Status.FAILED
            session.ErrorMessage = f"Превышен лимит строк ({MAX_ROWS})"
            db.commit()
            db.refresh(session)
            return session
        raw = {}
        for col in profile.columns:
            idx = column_map[col.field]
            raw[col.field] = row[idx] if idx is not None and idx < len(row) else None
        raw_rows.append((row_number, raw))

    normalized_rows = [
        (row_number, raw, *_normalize_row(db, profile, raw)) for row_number, raw in raw_rows
    ]

    key_counts: dict[tuple, list[int]] = {}
    for row_number, _raw, normalized, error in normalized_rows:
        if error or normalized is None:
            continue
        key = tuple(normalized.get(f) for f in profile.match_fields)
        if any(v is None for v in key):
            continue
        key_counts.setdefault(key, []).append(row_number)
    duplicate_keys = {k for k, nums in key_counts.items() if len(nums) > 1}

    for row_number, raw, normalized, error in normalized_rows:
        matched_id = None
        matched_version = None
        diff_fields = None
        validation_error = error
        classification = Classification.ERROR

        if not validation_error and normalized is not None:
            key = tuple(normalized.get(f) for f in profile.match_fields)
            if key in duplicate_keys:
                validation_error = "Дублирующийся ключ внутри файла"
            else:
                classification, matched_id, matched_version, diff_fields = _match_and_classify(
                    db, profile, normalized
                )

        db.add(
            ImportRow(
                ImportSessionId=session.Id,
                RowNumber=row_number,
                RawData=_json_safe_dict(raw),
                NormalizedData=_json_safe_dict(normalized),
                MatchedRecordId=matched_id,
                MatchedRecordVersion=matched_version,
                Classification=classification,
                DiffFields=diff_fields,
                ValidationError=validation_error,
            )
        )

    session.Status = Status.REVIEW
    db.commit()
    db.refresh(session)
    return session


def _deserialize_normalized(profile: ImportProfile, normalized: dict) -> dict:
    """Reverse of _json_safe for the fields a create/update schema expects."""
    result = {}
    for col in profile.columns:
        value = normalized.get(col.field)
        if value is None:
            result[col.field] = None
        elif col.kind == "date":
            result[col.field] = date.fromisoformat(value)
        elif col.kind == "enum":
            result[col.field] = col.enum_cls(value)
        elif col.kind == "fk_lookup":
            result[col.field] = uuid.UUID(value)
        else:
            result[col.field] = value
    return result


def revalidate_row(db: Session, profile: ImportProfile, row: ImportRow, raw_updates: dict) -> ImportRow:
    """Ручное исправление строки, попавшей в «Ошибки»: правки накатываются
    поверх сохранённой RawData, строка нормализуется и классифицируется
    заново — теми же правилами, что и при первичном разборе файла."""
    raw = _rehydrate_raw(profile, dict(row.RawData or {}))
    raw.update(raw_updates)
    normalized, error = _normalize_row(db, profile, raw)

    classification = Classification.ERROR
    matched_id = matched_version = diff_fields = None
    validation_error = error

    if not error and normalized is not None:
        key = tuple(normalized.get(f) for f in profile.match_fields)
        if any(v is None for v in key):
            validation_error = "Не заполнены поля сопоставления"
        elif _duplicate_within_session(db, row, profile, normalized):
            validation_error = "Дублирующийся ключ внутри файла"
        else:
            classification, matched_id, matched_version, diff_fields = _match_and_classify(
                db, profile, normalized
            )

    row.RawData = _json_safe_dict(raw)
    row.NormalizedData = _json_safe_dict(normalized) if normalized is not None else None
    row.MatchedRecordId = matched_id
    row.MatchedRecordVersion = matched_version
    row.Classification = classification if not validation_error else Classification.ERROR
    row.DiffFields = diff_fields
    row.ValidationError = validation_error
    row.Decision = Decision.PENDING
    row.AppliedAction = None
    row.AppliedError = None
    row.AppliedDate = None
    db.commit()
    db.refresh(row)
    return row


def _apply_row(db: Session, profile: ImportProfile, row: ImportRow, user_id: uuid.UUID | None) -> None:
    """Применяет одну подтверждённую строку к БД — общая логика для массового
    применения сессии и мгновенного применения при подтверждении строки."""
    try:
        if row.Classification == Classification.IDENTICAL:
            row.AppliedAction = AppliedAction.SKIPPED
        elif row.Classification == Classification.NEW:
            data = _deserialize_normalized(profile, row.NormalizedData)
            profile.repo.create(db, data, user_id)
            row.AppliedAction = AppliedAction.CREATED
        elif row.Classification == Classification.CONFLICT:
            current = profile.repo.get(db, row.MatchedRecordId)
            if current is None or current.Version != row.MatchedRecordVersion:
                row.AppliedAction = AppliedAction.FAILED
                row.AppliedError = (
                    "Запись изменена другим пользователем с момента сверки — требуется повторное ревью"
                )
            else:
                data = _deserialize_normalized(profile, row.NormalizedData)
                profile.repo.update(db, row.MatchedRecordId, data, user_id)
                row.AppliedAction = AppliedAction.UPDATED
        else:
            row.AppliedAction = AppliedAction.SKIPPED
        row.AppliedDate = func.now()
        db.commit()
    except Exception as exc:
        db.rollback()
        row.AppliedAction = AppliedAction.FAILED
        row.AppliedError = str(exc)
        db.commit()


def _finalize_session_status(db: Session, session: ImportSession, user_id: uuid.UUID | None = None) -> None:
    remaining_pending = db.scalar(
        select(func.count()).select_from(
            select(ImportRow.Id)
            .where(
                ImportRow.ImportSessionId == session.Id,
                ImportRow.Decision == Decision.CONFIRMED,
                ImportRow.AppliedAction.is_(None),
            )
            .subquery()
        )
    ) or 0
    any_failed = db.scalar(
        select(func.count()).select_from(
            select(ImportRow.Id)
            .where(
                ImportRow.ImportSessionId == session.Id,
                ImportRow.AppliedAction == AppliedAction.FAILED,
            )
            .subquery()
        )
    ) or 0
    any_applied = db.scalar(
        select(func.count()).select_from(
            select(ImportRow.Id)
            .where(
                ImportRow.ImportSessionId == session.Id,
                ImportRow.AppliedAction.isnot(None),
            )
            .subquery()
        )
    ) or 0

    if any_applied == 0:
        session.Status = Status.REVIEW
    elif any_failed or remaining_pending:
        session.Status = Status.PARTIALLY_APPLIED
    else:
        session.Status = Status.APPLIED
    if user_id is not None:
        session.AppliedBy = user_id
    session.AppliedDate = func.now()
    db.commit()


def set_row_decision(
    db: Session, session: ImportSession, row: ImportRow, decision: Decision, user_id: uuid.UUID | None
) -> ImportRow:
    if decision == Decision.CONFIRMED and row.Classification not in CONFIRMABLE_CLASSIFICATIONS:
        raise ValueError("Эту строку нельзя подтвердить (неоднозначно или ошибка)")
    row.Decision = decision
    db.commit()
    db.refresh(row)
    if decision == Decision.CONFIRMED:
        profile = get_profile(session.EntityType)
        _apply_row(db, profile, row, user_id)
        _finalize_session_status(db, session, user_id)
        db.refresh(row)
    return row


def apply_session(db: Session, session: ImportSession, user_id: uuid.UUID | None) -> dict:
    if session.Status not in (Status.REVIEW, Status.PARTIALLY_APPLIED, Status.APPLIED):
        raise ValueError("Сессия не в статусе, допускающем применение")

    profile = get_profile(session.EntityType)
    if profile is None:
        raise ValueError(f"Неизвестный тип импорта: {session.EntityType}")

    session.Status = Status.APPLYING
    db.commit()

    pending_rows = db.scalars(
        select(ImportRow).where(
            ImportRow.ImportSessionId == session.Id,
            ImportRow.Decision == Decision.CONFIRMED,
            ImportRow.AppliedAction.is_(None),
        )
    ).all()

    created = updated = failed = skipped = 0
    for row in pending_rows:
        _apply_row(db, profile, row, user_id)
        if row.AppliedAction == AppliedAction.CREATED:
            created += 1
        elif row.AppliedAction == AppliedAction.UPDATED:
            updated += 1
        elif row.AppliedAction == AppliedAction.FAILED:
            failed += 1
        else:
            skipped += 1

    remaining_pending = db.scalar(
        select(func.count()).select_from(
            select(ImportRow.Id)
            .where(
                ImportRow.ImportSessionId == session.Id,
                ImportRow.Decision == Decision.CONFIRMED,
                ImportRow.AppliedAction.is_(None),
            )
            .subquery()
        )
    ) or 0

    _finalize_session_status(db, session, user_id)

    log_action(
        db,
        table_name="ImportSessions",
        record_id=session.Id,
        action=AuditAction.IMPORT,
        user_id=user_id,
        new_value=f"created={created} updated={updated} failed={failed} skipped={skipped}",
    )
    db.commit()

    return {
        "created": created,
        "updated": updated,
        "failed": failed,
        "skipped": skipped,
        "remaining_pending": remaining_pending,
    }


def session_counts(db: Session, session_id: uuid.UUID) -> dict[str, int]:
    rows = db.scalars(
        select(ImportRow.Classification).where(ImportRow.ImportSessionId == session_id)
    ).all()
    counts = {c.value: 0 for c in Classification}
    for c in rows:
        counts[c.value] += 1
    return counts


def generate_template(profile: ImportProfile) -> bytes:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.append([col.header_aliases[0] for col in profile.columns])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
