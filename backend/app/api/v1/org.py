import uuid
from io import BytesIO
from typing import Annotated

import openpyxl
from fastapi import APIRouter, Depends, HTTPException, Response, UploadFile, status
from sqlalchemy import func, select

from app.api.deps import DbSession, require_permission
from app.api.v1.crud_router import build_crud_router
from app.models.employee import Employee
from app.models.enums import AuditAction, OrgUnitType
from app.models.org import City, LegalEntity, OrgUnit, Position
from app.models.user import User
from app.schemas.org import (
    AssignEmployeeRequest,
    CityCreate,
    CityOut,
    ImportSpecialistsResult,
    LegalEntityCreate,
    LegalEntityOut,
    OrgUnitChildItem,
    OrgUnitCreate,
    OrgUnitDetail,
    OrgUnitEmployeeItem,
    OrgUnitFullTreeItem,
    OrgUnitOut,
    OrgUnitTreeItem,
    OrgUnitUpdate,
    PositionCreate,
    PositionOut,
    SetManagerRequest,
)
from app.services.audit import log_action

# Жёсткий маршрут иерархии (клиент: "цепочка должна всегда следовать по
# этому маршруту") — Дивизион (верхний уровень, много) → Город (много на
# дивизион, во главе Директор филиала) → Управление (много на город) →
# Группа (много на управление) → Сотрудники. Пропуск уровня (например,
# Группа сразу под Дивизионом) запрещён на уровне API, а не только в UI.
ORG_UNIT_TYPE_LABELS = {
    OrgUnitType.DIVISION: "Дивизион",
    OrgUnitType.CITY: "Город",
    OrgUnitType.DEPARTMENT: "Управление",
    OrgUnitType.GROUP: "Группа",
}
ORG_UNIT_ALLOWED_PARENT_TYPES: dict[OrgUnitType, set[OrgUnitType]] = {
    OrgUnitType.DIVISION: set(),  # только верхний уровень, без родителя
    OrgUnitType.CITY: {OrgUnitType.DIVISION},
    OrgUnitType.DEPARTMENT: {OrgUnitType.CITY},
    OrgUnitType.GROUP: {OrgUnitType.DEPARTMENT},
}


def _validate_org_unit_hierarchy(db, unit_type: OrgUnitType, parent_id: uuid.UUID | None) -> None:
    allowed_parents = ORG_UNIT_ALLOWED_PARENT_TYPES[unit_type]
    label = ORG_UNIT_TYPE_LABELS[unit_type]

    if not allowed_parents:
        if parent_id is not None:
            raise HTTPException(
                status.HTTP_400_BAD_REQUEST,
                f"«{label}» — верхний уровень иерархии, у него не может быть родительского подразделения",
            )
        return

    if parent_id is None:
        allowed_labels = ", ".join(ORG_UNIT_TYPE_LABELS[t] for t in allowed_parents)
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            f"«{label}» обязательно должен быть внутри одного из: {allowed_labels}",
        )

    parent = db.get(OrgUnit, parent_id)
    if parent is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Родительское подразделение не найдено")
    if parent.UnitType not in allowed_parents:
        allowed_labels = ", ".join(ORG_UNIT_TYPE_LABELS[t] for t in allowed_parents)
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            f"«{label}» может находиться только внутри: {allowed_labels} "
            f"(выбранное подразделение имеет тип «{ORG_UNIT_TYPE_LABELS.get(parent.UnitType, parent.UnitType)}»)",
        )


def _validate_org_unit_create(db, payload: dict) -> None:
    _validate_org_unit_hierarchy(db, payload["UnitType"], payload.get("ParentOrgUnitId"))


def _validate_org_unit_update(db, existing: OrgUnit, payload: dict) -> None:
    # PATCH может прислать только часть полей — проверяем итоговое состояние
    # (новое значение, если пришло, иначе то, что уже сохранено в записи).
    unit_type = payload.get("UnitType", existing.UnitType)
    parent_id = payload["ParentOrgUnitId"] if "ParentOrgUnitId" in payload else existing.ParentOrgUnitId
    _validate_org_unit_hierarchy(db, unit_type, parent_id)


org_units_router = build_crud_router(
    model=OrgUnit,
    create_schema=OrgUnitCreate,
    update_schema=OrgUnitUpdate,
    out_schema=OrgUnitOut,
    prefix="/org-units",
    tags=["org-units"],
    search_fields=["Name"],
    module="ORG_UNITS",
    validate_create=_validate_org_unit_create,
    validate_update=_validate_org_unit_update,
)

org_units_extra_router = APIRouter(prefix="/org-units", tags=["org-units"])

ViewOrgUnits = Annotated[User, Depends(require_permission("ORG_UNITS", "VIEW"))]
UpdateOrgUnits = Annotated[User, Depends(require_permission("ORG_UNITS", "UPDATE"))]
ExportOrgUnits = Annotated[User, Depends(require_permission("ORG_UNITS", "EXPORT"))]
ImportOrgUnits = Annotated[User, Depends(require_permission("ORG_UNITS", "IMPORT"))]


def _get_org_unit_or_404(db, org_unit_id: uuid.UUID) -> OrgUnit:
    unit = db.get(OrgUnit, org_unit_id)
    if unit is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Подразделение не найдено")
    return unit


@org_units_extra_router.get("/tree", response_model=list[OrgUnitTreeItem])
def get_org_unit_tree(db: DbSession, _user: ViewOrgUnits):
    counts = dict(
        db.execute(
            select(Employee.OrgUnitId, func.count(Employee.Id))
            .where(Employee.OrgUnitId.is_not(None), Employee.IsActive.is_(True))
            .group_by(Employee.OrgUnitId)
        ).all()
    )
    units = db.scalars(select(OrgUnit).where(OrgUnit.IsActive.is_(True))).all()
    return [
        OrgUnitTreeItem(
            Id=u.Id,
            Name=u.Name,
            UnitType=u.UnitType,
            ParentOrgUnitId=u.ParentOrgUnitId,
            IsActive=u.IsActive,
            EmployeeCount=counts.get(u.Id, 0),
        )
        for u in units
    ]


@org_units_extra_router.get("/full-tree", response_model=list[OrgUnitFullTreeItem])
def get_org_unit_full_tree(db: DbSession, _user: ViewOrgUnits):
    """Наглядный реестр «кто за кем идёт» (руководитель каждого уровня +
    сотрудники Групп) одним запросом — питает вкладку «Список», где
    N обращений к /detail на узел были бы неприемлемо медленными."""
    units = db.scalars(select(OrgUnit).where(OrgUnit.IsActive.is_(True))).all()
    employees = db.scalars(select(Employee).where(Employee.IsActive.is_(True))).all()
    positions = {p.Id: p.Name for p in db.scalars(select(Position))}

    employees_by_id = {e.Id: e for e in employees}
    employees_by_unit: dict[uuid.UUID, list[Employee]] = {}
    for e in employees:
        if e.OrgUnitId:
            employees_by_unit.setdefault(e.OrgUnitId, []).append(e)

    children_by_parent: dict[uuid.UUID | None, list[OrgUnit]] = {}
    for u in units:
        children_by_parent.setdefault(u.ParentOrgUnitId, []).append(u)

    def build(u: OrgUnit) -> OrgUnitFullTreeItem:
        manager = employees_by_id.get(u.ManagerEmployeeId) if u.ManagerEmployeeId else None
        return OrgUnitFullTreeItem(
            Id=u.Id,
            Name=u.Name,
            UnitType=u.UnitType,
            ManagerName=manager.FullName if manager else None,
            Employees=[
                OrgUnitEmployeeItem(
                    Id=e.Id,
                    FullName=e.FullName,
                    PositionName=positions.get(e.PositionId),
                    Status=e.Status.value,
                )
                for e in employees_by_unit.get(u.Id, [])
            ],
            Children=[build(c) for c in children_by_parent.get(u.Id, [])],
        )

    return [build(u) for u in children_by_parent.get(None, [])]


_ORG_UNIT_LEVEL_ORDER = {t: i for i, t in enumerate(ORG_UNIT_TYPE_LABELS)}


@org_units_extra_router.get("/export")
def export_org_units(db: DbSession, user: ExportOrgUnits):
    """Экспорт всей оргструктуры в Excel — для передачи данных другой
    стороне и для использования в дальнейших расчётах вне системы."""
    units = db.scalars(select(OrgUnit).where(OrgUnit.IsActive.is_(True))).all()
    names_by_id = {u.Id: u.Name for u in units}
    manager_ids = {u.ManagerEmployeeId for u in units if u.ManagerEmployeeId}
    manager_names = {
        e.Id: e.FullName
        for e in db.scalars(select(Employee).where(Employee.Id.in_(manager_ids)))
    } if manager_ids else {}

    units_sorted = sorted(
        units, key=lambda u: (_ORG_UNIT_LEVEL_ORDER.get(u.UnitType, 99), u.Name)
    )

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Оргструктура"
    sheet.append(["Тип", "Название", "Код", "Родитель", "Руководитель"])
    for u in units_sorted:
        sheet.append([
            ORG_UNIT_TYPE_LABELS.get(u.UnitType, u.UnitType),
            u.Name,
            u.Code or "",
            names_by_id.get(u.ParentOrgUnitId, ""),
            manager_names.get(u.ManagerEmployeeId, ""),
        ])

    log_action(
        db, table_name="OrgUnits", record_id=uuid.uuid4(), action=AuditAction.EXPORT,
        user_id=user.Id, reason="Экспорт организационной структуры",
    )
    db.commit()

    buffer = BytesIO()
    workbook.save(buffer)
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=org_units_export.xlsx"},
    )


SPECIALIST_EXPORT_HEADERS = [
    "Название", "Руководитель дивизиона", "Город", "ДФ",
    "Управление", "РУ", "РГ", "Группа",
]


@org_units_extra_router.get("/export-specialists")
def export_specialists(db: DbSession, user: ExportOrgUnits):
    """Экспорт специалистов построчно с полной цепочкой руководителей — тот
    же формат принимает /import-specialists для массовой загрузки обратно."""
    units = db.scalars(select(OrgUnit).where(OrgUnit.IsActive.is_(True))).all()
    units_by_id = {u.Id: u for u in units}
    employees = db.scalars(select(Employee).where(Employee.IsActive.is_(True))).all()
    employees_by_id = {e.Id: e for e in employees}

    def manager_name(unit: OrgUnit | None) -> str:
        if unit is None or not unit.ManagerEmployeeId:
            return ""
        manager = employees_by_id.get(unit.ManagerEmployeeId)
        return manager.FullName if manager else ""

    rows = []
    for emp in employees:
        group = units_by_id.get(emp.OrgUnitId) if emp.OrgUnitId else None
        if group is None or group.UnitType != OrgUnitType.GROUP:
            continue
        department = units_by_id.get(group.ParentOrgUnitId) if group.ParentOrgUnitId else None
        city = units_by_id.get(department.ParentOrgUnitId) if department and department.ParentOrgUnitId else None
        division = units_by_id.get(city.ParentOrgUnitId) if city and city.ParentOrgUnitId else None
        rows.append((
            division.Name if division else "",
            manager_name(division),
            city.Name if city else "",
            manager_name(city),
            department.Name if department else "",
            manager_name(department),
            manager_name(group),
            emp.FullName,
        ))

    rows.sort(key=lambda r: (r[0], r[2], r[4], r[6], r[7]))

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Специалисты"
    sheet.append(SPECIALIST_EXPORT_HEADERS)
    for r in rows:
        sheet.append(list(r))

    log_action(
        db, table_name="OrgUnits", record_id=uuid.uuid4(), action=AuditAction.EXPORT,
        user_id=user.Id, reason="Экспорт специалистов по подразделениям",
    )
    db.commit()

    buffer = BytesIO()
    workbook.save(buffer)
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=specialists_export.xlsx"},
    )


_SPECIALIST_IMPORT_COLUMNS = {
    "название": "division",
    "город": "city",
    "управление": "department",
    "рг": "group_manager",
    "группа": "specialist",
}


@org_units_extra_router.post("/import-specialists", response_model=ImportSpecialistsResult)
async def import_specialists(db: DbSession, user: ImportOrgUnits, file: UploadFile):
    """Массовая загрузка специалистов по файлу в формате /export-specialists:
    подразделения (Дивизион/Город/Управление) должны уже существовать —
    файл только сопоставляет и назначает существующих сотрудников их Группе,
    не создаёт структуру и не создаёт сотрудников (см. решение по Импорту
    OrgUnit из истории этой сессии — та же логика "сопоставление, не
    создание", просто на уровне назначений, а не самих подразделений)."""
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Поддерживаются только файлы .xlsx")

    content = await file.read()
    try:
        workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows_iter = sheet.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
    except Exception:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST, "Не удалось прочитать файл — убедитесь, что это корректный .xlsx"
        )
    if header_row is None:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Файл пуст")

    col_index: dict[str, int] = {}
    for i, h in enumerate(header_row):
        key = str(h).strip().lower() if h is not None else None
        if key in _SPECIALIST_IMPORT_COLUMNS:
            col_index[_SPECIALIST_IMPORT_COLUMNS[key]] = i
    missing = [
        label for key, label in [
            ("division", "Название"), ("city", "Город"),
            ("department", "Управление"), ("specialist", "Группа"),
        ]
        if key not in col_index
    ]
    if missing:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            f"В файле не найдены обязательные колонки: {', '.join(missing)}",
        )

    def cell(row: tuple, key: str) -> str | None:
        idx = col_index.get(key)
        if idx is None or idx >= len(row):
            return None
        value = row[idx]
        return str(value).strip() if value is not None else None

    assigned = unchanged = failed = 0
    errors = []

    for row_number, row in enumerate(rows_iter, start=2):
        if row is None or all(v is None for v in row):
            continue

        division_name = cell(row, "division")
        city_name = cell(row, "city")
        department_name = cell(row, "department")
        group_manager_name = cell(row, "group_manager")
        specialist_name = cell(row, "specialist")

        if not division_name or not city_name or not department_name or not specialist_name:
            failed += 1
            errors.append({
                "RowNumber": row_number,
                "Message": "Не заполнены обязательные поля (Название/Город/Управление/Группа)",
            })
            continue

        division = db.scalar(
            select(OrgUnit).where(
                OrgUnit.Name == division_name, OrgUnit.UnitType == OrgUnitType.DIVISION,
                OrgUnit.IsActive.is_(True),
            )
        )
        if division is None:
            failed += 1
            errors.append({"RowNumber": row_number, "Message": f"Дивизион не найден: «{division_name}»"})
            continue

        city = db.scalar(
            select(OrgUnit).where(
                OrgUnit.Name == city_name, OrgUnit.UnitType == OrgUnitType.CITY,
                OrgUnit.ParentOrgUnitId == division.Id, OrgUnit.IsActive.is_(True),
            )
        )
        if city is None:
            failed += 1
            errors.append({
                "RowNumber": row_number,
                "Message": f"Город «{city_name}» не найден в дивизионе «{division_name}»",
            })
            continue

        department = db.scalar(
            select(OrgUnit).where(
                OrgUnit.Name == department_name, OrgUnit.UnitType == OrgUnitType.DEPARTMENT,
                OrgUnit.ParentOrgUnitId == city.Id, OrgUnit.IsActive.is_(True),
            )
        )
        if department is None:
            failed += 1
            errors.append({
                "RowNumber": row_number,
                "Message": f"Управление «{department_name}» не найдено в городе «{city_name}»",
            })
            continue

        groups = db.scalars(
            select(OrgUnit).where(
                OrgUnit.UnitType == OrgUnitType.GROUP, OrgUnit.ParentOrgUnitId == department.Id,
                OrgUnit.IsActive.is_(True),
            )
        ).all()
        if not groups:
            failed += 1
            errors.append({"RowNumber": row_number, "Message": f"В «{department_name}» нет ни одной группы"})
            continue

        group = None
        if len(groups) == 1 and not group_manager_name:
            group = groups[0]
        else:
            manager_ids = {g.ManagerEmployeeId for g in groups if g.ManagerEmployeeId}
            managers = {
                e.Id: e for e in db.scalars(select(Employee).where(Employee.Id.in_(manager_ids)))
            } if manager_ids else {}
            matches = [
                g for g in groups
                if g.ManagerEmployeeId
                and managers.get(g.ManagerEmployeeId)
                and managers[g.ManagerEmployeeId].FullName.strip().lower()
                == (group_manager_name or "").strip().lower()
            ]
            if len(matches) == 1:
                group = matches[0]
        if group is None:
            failed += 1
            errors.append({
                "RowNumber": row_number,
                "Message": (
                    f"Не удалось определить группу в «{department_name}» "
                    f"по руководителю «{group_manager_name or '—'}»"
                ),
            })
            continue

        matching_employees = db.scalars(
            select(Employee).where(Employee.FullName == specialist_name, Employee.IsActive.is_(True))
        ).all()
        if len(matching_employees) == 0:
            failed += 1
            errors.append({"RowNumber": row_number, "Message": f"Сотрудник не найден: «{specialist_name}»"})
            continue
        if len(matching_employees) > 1:
            failed += 1
            errors.append({
                "RowNumber": row_number,
                "Message": f"Найдено несколько сотрудников с ФИО «{specialist_name}» — уточните вручную",
            })
            continue
        employee = matching_employees[0]

        if employee.OrgUnitId == group.Id:
            unchanged += 1
            continue

        old_org_unit_id = employee.OrgUnitId
        log_action(
            db, table_name="Employees", record_id=employee.Id, action=AuditAction.UPDATE,
            user_id=user.Id, field_name="OrgUnitId",
            old_value=str(old_org_unit_id) if old_org_unit_id else None, new_value=str(group.Id),
        )
        employee.OrgUnitId = group.Id
        employee.ModifiedBy = user.Id
        employee.Version += 1
        assigned += 1

    db.commit()
    return ImportSpecialistsResult(Assigned=assigned, Unchanged=unchanged, Failed=failed, Errors=errors)


@org_units_extra_router.get("/{org_unit_id}/detail", response_model=OrgUnitDetail)
def get_org_unit_detail(org_unit_id: uuid.UUID, db: DbSession, _user: ViewOrgUnits):
    unit = _get_org_unit_or_404(db, org_unit_id)

    manager_name = None
    if unit.ManagerEmployeeId:
        manager = db.get(Employee, unit.ManagerEmployeeId)
        manager_name = manager.FullName if manager else None

    parent_name = None
    if unit.ParentOrgUnitId:
        parent = db.get(OrgUnit, unit.ParentOrgUnitId)
        parent_name = parent.Name if parent else None

    children = db.scalars(
        select(OrgUnit).where(
            OrgUnit.ParentOrgUnitId == org_unit_id, OrgUnit.IsActive.is_(True)
        )
    ).all()

    employees = db.scalars(
        select(Employee).where(
            Employee.OrgUnitId == org_unit_id, Employee.IsActive.is_(True)
        )
    ).all()
    position_ids = {e.PositionId for e in employees if e.PositionId}
    positions = {
        p.Id: p.Name
        for p in db.scalars(select(Position).where(Position.Id.in_(position_ids)))
    } if position_ids else {}

    return OrgUnitDetail(
        Id=unit.Id,
        Name=unit.Name,
        UnitType=unit.UnitType,
        Code=unit.Code,
        ParentOrgUnitId=unit.ParentOrgUnitId,
        ManagerEmployeeId=unit.ManagerEmployeeId,
        CreatedDate=unit.CreatedDate,
        CreatedBy=unit.CreatedBy,
        ModifiedDate=unit.ModifiedDate,
        ModifiedBy=unit.ModifiedBy,
        IsActive=unit.IsActive,
        Version=unit.Version,
        Comment=unit.Comment,
        ManagerName=manager_name,
        ParentName=parent_name,
        Children=[
            OrgUnitChildItem(Id=c.Id, Name=c.Name, UnitType=c.UnitType) for c in children
        ],
        Employees=[
            OrgUnitEmployeeItem(
                Id=e.Id,
                FullName=e.FullName,
                PositionName=positions.get(e.PositionId),
                Status=e.Status.value,
            )
            for e in employees
        ],
    )


@org_units_extra_router.post("/{org_unit_id}/assign-employee", response_model=OrgUnitEmployeeItem)
def assign_employee(
    org_unit_id: uuid.UUID, data: AssignEmployeeRequest, db: DbSession, user: UpdateOrgUnits
):
    unit = _get_org_unit_or_404(db, org_unit_id)
    if unit.UnitType != OrgUnitType.GROUP:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            "Сотрудников можно закрепить только за «Группой» — нижним уровнем иерархии",
        )
    employee = db.get(Employee, data.EmployeeId)
    if employee is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Сотрудник не найден")

    old_org_unit_id = employee.OrgUnitId
    log_action(
        db,
        table_name="Employees",
        record_id=employee.Id,
        action=AuditAction.UPDATE,
        user_id=user.Id,
        field_name="OrgUnitId",
        old_value=str(old_org_unit_id) if old_org_unit_id else None,
        new_value=str(org_unit_id),
    )
    employee.OrgUnitId = org_unit_id
    employee.ModifiedBy = user.Id
    employee.Version += 1
    db.commit()

    position_name = None
    if employee.PositionId:
        position = db.get(Position, employee.PositionId)
        position_name = position.Name if position else None
    return OrgUnitEmployeeItem(
        Id=employee.Id,
        FullName=employee.FullName,
        PositionName=position_name,
        Status=employee.Status.value,
    )


@org_units_extra_router.post("/{org_unit_id}/remove-employee", status_code=status.HTTP_204_NO_CONTENT)
def remove_employee(
    org_unit_id: uuid.UUID, data: AssignEmployeeRequest, db: DbSession, user: UpdateOrgUnits
):
    _get_org_unit_or_404(db, org_unit_id)
    employee = db.get(Employee, data.EmployeeId)
    if employee is None or employee.OrgUnitId != org_unit_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Сотрудник не числится в этом подразделении")

    log_action(
        db,
        table_name="Employees",
        record_id=employee.Id,
        action=AuditAction.UPDATE,
        user_id=user.Id,
        field_name="OrgUnitId",
        old_value=str(org_unit_id),
        new_value=None,
    )
    employee.OrgUnitId = None
    employee.ModifiedBy = user.Id
    employee.Version += 1
    db.commit()


@org_units_extra_router.post("/{org_unit_id}/set-manager", response_model=OrgUnitOut)
def set_manager(
    org_unit_id: uuid.UUID, data: SetManagerRequest, db: DbSession, user: UpdateOrgUnits
):
    unit = _get_org_unit_or_404(db, org_unit_id)
    if data.EmployeeId is not None and db.get(Employee, data.EmployeeId) is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Сотрудник не найден")

    if data.EmployeeId is not None and not data.Force:
        conflict = db.scalar(
            select(OrgUnit).where(
                OrgUnit.ManagerEmployeeId == data.EmployeeId,
                OrgUnit.Id != org_unit_id,
                OrgUnit.IsActive.is_(True),
            )
        )
        if conflict is not None:
            conflict_label = ORG_UNIT_TYPE_LABELS.get(conflict.UnitType, conflict.UnitType)
            raise HTTPException(
                status.HTTP_409_CONFLICT,
                detail={
                    "conflict": True,
                    "message": (
                        f"Этот сотрудник уже назначен «{conflict_label}» в «{conflict.Name}». "
                        "Всё равно назначить здесь?"
                    ),
                    "conflictUnitId": str(conflict.Id),
                    "conflictUnitName": conflict.Name,
                },
            )

    log_action(
        db,
        table_name="OrgUnits",
        record_id=unit.Id,
        action=AuditAction.UPDATE,
        user_id=user.Id,
        field_name="ManagerEmployeeId",
        old_value=str(unit.ManagerEmployeeId) if unit.ManagerEmployeeId else None,
        new_value=str(data.EmployeeId) if data.EmployeeId else None,
    )
    unit.ManagerEmployeeId = data.EmployeeId
    unit.ModifiedBy = user.Id
    unit.Version += 1
    db.commit()
    db.refresh(unit)
    return unit


positions_router = build_crud_router(
    model=Position,
    create_schema=PositionCreate,
    update_schema=PositionCreate,
    out_schema=PositionOut,
    prefix="/positions",
    tags=["positions"],
    search_fields=["Name"],
    module="POSITIONS",
)

cities_router = build_crud_router(
    model=City,
    create_schema=CityCreate,
    update_schema=CityCreate,
    out_schema=CityOut,
    prefix="/cities",
    tags=["cities"],
    search_fields=["Name"],
    module="CITIES",
)

legal_entities_router = build_crud_router(
    model=LegalEntity,
    create_schema=LegalEntityCreate,
    update_schema=LegalEntityCreate,
    out_schema=LegalEntityOut,
    prefix="/legal-entities",
    tags=["legal-entities"],
    search_fields=["Name"],
    module="LEGAL_ENTITIES",
)
